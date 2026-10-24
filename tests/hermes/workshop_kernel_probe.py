"""Runs as the real offline task UID, after production namespace and seccomp setup."""
import asyncio
import ctypes
import errno
import fcntl
import json
import os
from pathlib import Path
import socket
import subprocess
import sys

expected = json.loads(sys.argv[1])
libc = ctypes.CDLL(None, use_errno=True)
assert os.geteuid() >= 20001
for namespace in ('net', 'ipc', 'mnt', 'pid'):
    assert os.readlink('/proc/self/ns/' + namespace) != expected['namespaces'][namespace]
assert not Path('/tmp/' + expected['marker']).exists()
assert not Path('/dev/shm/' + expected['marker']).exists()
assert not Path('/dev/mqueue' + expected['mq']).exists()
assert not Path('/workshop/socket').exists()
assert not Path('/workshop/fifo').exists()
file = Path('/workshop/document.txt')
assert file.read_text() == 'synthetic document\n'
assert [file.stat().st_dev, file.stat().st_ino] != expected['inode']
with file.open() as stream:
    fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
try:
    file.write_text('changed')
except OSError as error:
    assert error.errno == errno.EROFS
else:
    raise AssertionError('snapshot mount is writable')
assert libc.msgget(expected['queue_key'], 0) == -1 and ctypes.get_errno() == errno.ENOENT
assert libc.mq_open(expected['mq'].encode(), os.O_RDWR) == -1
assert ctypes.get_errno() == errno.ENOENT

async def event_loop():
    await asyncio.sleep(0)
    return 'asyncio-ok'

assert asyncio.run(event_loop()) == 'asyncio-ok'
node = subprocess.run(['node', '-e', "const {spawn}=require('child_process');const p=spawn('/bin/cat',[],{stdio:'pipe'});let s='';p.stdout.on('data',b=>s+=b);p.on('close',c=>{if(c||s!=='pipe-ok')process.exit(1);console.log('node-ok')});p.stdin.end('pipe-ok')"],
                      capture_output=True, text=True, check=True)
assert node.stdout.strip() == 'node-ok'
address = '/tmp/local-test.sock'
with socket.socket(socket.AF_UNIX) as server, socket.socket(socket.AF_UNIX) as client:
    server.bind(address)
    server.listen()
    client.connect(address)
    connection, _ = server.accept()
    with connection:
        client.sendall(b'local-ok')
        assert connection.recv(8) == b'local-ok'
with socket.socket() as client:
    assert client.connect_ex(('127.0.0.1', expected['port'])) != 0
from docx import Document
document = Document()
document.add_paragraph(file.read_text())
output = Path(os.environ['OUT']) / 'offline.docx'
document.save(output)
assert Document(output).paragraphs[0].text == 'synthetic document\n'
# Resolve the architecture-specific syscall number through the installed libseccomp.
seccomp = ctypes.CDLL('libseccomp.so.2')
seccomp.seccomp_syscall_resolve_name.argtypes = [ctypes.c_char_p]
number = seccomp.seccomp_syscall_resolve_name(b'io_uring_setup')
assert number >= 0
assert libc.syscall(number, 1, ctypes.create_string_buffer(256)) == -1
assert ctypes.get_errno() == errno.EPERM
print(json.dumps({'offline_kernel_proof': 'passed', 'uid': os.geteuid(),
                  'document_bytes': output.stat().st_size}), flush=True)
